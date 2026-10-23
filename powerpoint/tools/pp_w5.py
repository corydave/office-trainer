# -*- coding: utf-8 -*-
"""
Week 5 practice presentations: Present and Collaborate, and the capstone practice exam.

build(path) makes the starter; build(path, done={...}) makes it with those
steps finished, for testing the checks. Everything is fictional.
"""
import os
from pptkit import *
from pptx.enum.shapes import MSO_SHAPE
from pp_w3 import add_chart, ICON_NAMES, _icons_row
from pp_w4 import _link_text, SITE

ORG = "Lakeshore Paddle Company"
SAM, ANA = "Sam Lee", "Ana Ortiz"
ME = "Jordan Rivera"   # stands in for the learner in the test files


# ---------------------------------------------------------------------------
# Lesson 1: Comments and Co-authoring  (D5.01 comments; D5.02 co-authoring and version history)
# ---------------------------------------------------------------------------
def build_l1(path, done=frozenset()):
    slides = [
        dict(layout="Title Slide", title="Summer events", subtitle=ORG),
        dict(title="Sunset tours", body=["Fridays at 7 p.m.", "Guide and snacks included"]),
        dict(title="Hours", body=["Open daily, 9 a.m. to 7 p.m."]),
        dict(title="Staff picnic", body=["Saturday, June 12, at noon"]),
        dict(title="Rain policy", body=["Thunder: everyone off the water", "Full refund if we cancel"]),
    ]
    if "partner" in done:
        slides.insert(4, dict(title="Partner's ideas", body=["A family paddle day in August"]))
    if "deleted" in done:   # a slide deleted by mistake, not restored yet
        slides = [s for s in slides if s["title"] != "Rain policy"]

    def after(p):
        sam = {"slide": 2, "author": SAM, "text": "Can we add the sunset tour price?"}
        if "reply" in done:
            sam["replies"] = [(ME, "Yes: $35 a person.")]
        if "resolve" in done:
            sam["resolved"] = True
        threads = [sam]
        if "delete" not in done:
            threads.append({"slide": 3, "author": ANA, "text": "Old: check last year's hours"})
        if "add" in done:
            threads.append({"slide": 4, "author": ME, "text": "Who brings the grill?"})
        comments(p, threads)
    return build_deck(path, slides, title="Summer events", author=ORG, after_save=after)


# ---------------------------------------------------------------------------
# Lesson 2: Compare and Combine  (D5.03)
# ---------------------------------------------------------------------------
def build_l2(path, done=frozenset()):
    """Your copy. done: 'reviewer' makes the reviewer's copy (all three changes);
    'title', 'price', 'color' are the reviewer's changes accepted into yours."""
    rev = "reviewer" in done

    def hours(prs, s):
        if rev or "color" in done:
            set_text_color([p for p in s.placeholders if p.placeholder_format.idx == 1][0], (192, 0, 0))

    slides = [
        dict(layout="Title Slide", title="Summer kickoff 2027" if (rev or "title" in done) else "Summer kickoff", subtitle=ORG),
        dict(title="Prices", body=["Kayak: $22 an hour" if (rev or "price" in done) else "Kayak: $20 an hour", "Canoe: $25 an hour"]),
        dict(title="Hours", body=["Open daily, 9 a.m. to 7 p.m.", "Last rental at 5 p.m."], fn=hours),
        dict(title="Thank you", body=["See you on the water"]),
    ]
    return build_deck(path, slides, title="Summer kickoff", author=SAM if rev else ORG)


# ---------------------------------------------------------------------------
# Lesson 3: Deliver the Show  (D5.04 custom show and show options; D5.05 timings, recording options, Presenter View)
# ---------------------------------------------------------------------------
TOUR = ["Lakeshore tour", "Welcome", "Our boats", "Prices", "Safety", "Weather rules", "Book online", "Thank you"]


def build_l3(path, done=frozenset()):
    slides = [dict(layout="Title Slide", title=TOUR[0], subtitle=ORG)] + [dict(title=t, body=["About " + t.lower()]) for t in TOUR[1:]]

    def before(prs):
        if "rehearse" in done:
            for i, s in enumerate(prs.slides):
                transition(s, "cut", after=6 + i % 3)

    def after(p):
        show_settings(p, loop="loop" in done or "kiosk" in done, kiosk="kiosk" in done,
                      narration="recording" in done, use_timings="recording" in done)
        if "custom" in done:
            custom_show(p, "Short tour", [2, 3, 7, 8])
    return build_deck(path, slides, title="Lakeshore tour", before_save=before, after_save=after)


# ---------------------------------------------------------------------------
# Lesson 4: Inspect, Protect and Share  (D5.06; D5.07 export)
# ---------------------------------------------------------------------------
def build_l4(path, done=frozenset()):
    def photo(prs, s):
        pic = picture(s, media("lake2.jpg"), 1, 1.6, 6, name="Picture 3")
        if "alt" in done:
            set_alt(pic, "Kayaks lined up on the beach at sunrise")
        if "offslide" not in done:
            text_box(s, 14, 2, 3, 1, ["TODO: check prices with Sam"], size=18, name="Note to self")

    def prices(prs, s):
        rows = [["Boat", "Hour", "Day"], ["Kayak", "$20", "$70"], ["Canoe", "$25", "$85"]]
        t = s.shapes.add_table(3, 3, Inches(1), Inches(1.8), Inches(8), Inches(1.5))
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                t.table.cell(i, j).text = v
        if "header" not in done:
            t.table.first_row = False

    slides = [
        dict(layout="Title Slide", title="Season wrap-up", subtitle=ORG),
        dict(layout="Title Only", title="Our best summer", fn=photo, notes="Thank everyone by name."),
        dict(layout="Title Only", title="Prices", fn=prices),
        dict(layout="Title Only", title="Thank you" if "slidetitle" in done else ""),
    ]

    def after(p):
        if "inspect" not in done:
            comments(p, [{"slide": 2, "author": SAM, "text": "Is this the right photo?"}])
        pres_flags(p, embed_fonts="fonts" in done, remove_personal="inspect" in done, modify_password="password" in done)
        if "final" in done:
            mark_final(p)
    return build_deck(path, slides, title="Season wrap-up", author=ME, after_save=after)


# ---------------------------------------------------------------------------
# Skill Check 5: team presentation
# ---------------------------------------------------------------------------
def build_sc5(path, done=frozenset()):
    def photo(prs, s):
        pic = picture(s, media("lake3.jpg"), 1, 1.6, 6, name="Picture 3")
        if "alt" in done:
            set_alt(pic, "The team paddling together on a calm morning")
        if "offslide" not in done:
            text_box(s, -4, 2, 3, 1, ["Old caption"], size=18, name="Old caption")

    titles = ["Team retreat", "Agenda", "Our morning paddle", "Lunch", "Afternoon games", "Thank you"]
    slides = [dict(layout="Title Slide", title=titles[0], subtitle=ORG)]
    for t in titles[1:]:
        slides.append(dict(layout="Title Only", title=t, fn=photo) if t == "Our morning paddle" else dict(title=t, body=["About " + t.lower()]))

    def before(prs):
        if "rehearse" in done:
            for s in prs.slides:
                transition(s, "cut", after=10)

    def after(p):
        q = {"slide": 2, "author": SAM, "text": "Is the date right?"}
        if "reply" in done:
            q["replies"] = [(ME, "Yes, June 12.")]
        if "resolve" in done:
            q["resolved"] = True
        threads = [q]
        if "delete" not in done:
            threads.append({"slide": 4, "author": ANA, "text": "Pizza or sandwiches?"})
        if "add" in done:
            threads.append({"slide": 6, "author": ME, "text": "Add our email?"})
        comments(p, threads)
        if "custom" in done:
            custom_show(p, "Highlights", [1, 3, 6])
        if "loop" in done:
            show_settings(p, loop=True)
        pres_flags(p, embed_fonts="fonts" in done, remove_personal="personal" in done)
    return build_deck(path, slides, title="Team retreat", author=ME, before_save=before, after_save=after)


# ---------------------------------------------------------------------------
# Capstone practice exam
# ---------------------------------------------------------------------------
def build_cap(path, done=frozenset()):
    refs = {}

    def highlights(prs, s):
        refs["list"] = [p for p in s.placeholders if p.placeholder_format.idx == 1][0]
        if "shape" in done:
            autoshape(s, MSO_SHAPE.EXPLOSION1, 9, 4, 3.4, 2, "Record year!")

    def beach(prs, s):
        pic = picture(s, media("lake4.jpg"), 1, 1.6, 6, name="Beach photo")
        if "style" in done:
            picture_style(pic)
        if "alt" in done:
            set_alt(pic, "Boats and families on the beach in July")

    def icons_slide(prs, s):
        _icons_row(s, {"align"} if "align" in done else set(), top=2.6)

    def prices(prs, s):
        rows = [["Boat", "Hour", "Day"], ["Kayak", "$20", "$70"], ["Canoe", "$25", "$85"]]
        if "row" in done:
            rows.append(["Paddleboard", "$18", "$60"])
        t = s.shapes.add_table(len(rows), 3, Inches(1), Inches(1.8), Inches(8), Inches(0.5 * len(rows)))
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                t.table.cell(i, j).text = v
        if "tablestyle" in done:
            table_style(t, MEDIUM_2_ACCENT_2)

    def rentals(prs, s):
        add_chart(s, "line" if "charttype" in done else "column", "Rentals by month")

    def safety(prs, s):
        tb = text_box(s, 0.8, 1.6, 9, 4, ["Check the weather", "Put on a life jacket", "Tell someone your plan"], size=24, name="Safety steps")
        if "numbers" in done:
            bullets(tb, "number")

    def footers(prs, s):
        if "footer" in done:
            footer_placeholders(prs, s, None, number=True)

    slides = [
        dict(layout="Title Slide", title="Lakeshore season review", subtitle=ORG),
        dict(title="Highlights", body=["Busiest July ever", "Two new guides", "Five-star reviews"], fn=lambda p, s: (highlights(p, s), footers(p, s))),
        dict(layout="Title Only", title="Our beach", fn=lambda p, s: (beach(p, s), footers(p, s))),
        dict(layout="Title Only", title="What we offer", fn=lambda p, s: (icons_slide(p, s), footers(p, s))),
        dict(layout="Title Only", title="Prices", fn=lambda p, s: (prices(p, s), footers(p, s))),
        dict(layout="Title Only", title="Rentals", fn=lambda p, s: (rentals(p, s), footers(p, s))),
        dict(title="Draft numbers", body=["Not final"], hidden="hide" in done, fn=footers),
        dict(layout="Title Only", title="Safety", fn=lambda p, s: (safety(p, s), footers(p, s))),
        dict(title="Our process", body=None if "smartart" in done else ["Reserve", "Paddle", "Return"], fn=footers),
        dict(title="Visit us", body=["Find us on our website"], fn=footers),
    ]
    if "newslide" in done:
        slides.insert(2, dict(title="Next steps", body=["Hire a third guide"], fn=footers))
    secs = None
    if "section" in done:
        secs = [("Default Section", "Lakeshore season review"), ("Results", "Rentals")]

    def before(prs):
        S = list(prs.slides)
        if "fade" in done:
            for s in S:
                transition(s, "fade")
        if "anim" in done:
            hl = [s for s in S if s.shapes.title is not None and s.shapes.title.text == "Highlights"][0]
            animate(hl, [dict(shape=refs["list"], cls="entr", preset=22, subtype=4, paragraphs=True, dur=1)])
        if "link" in done:
            _link_text(S[-1], url="https://www.lakeshorepaddle.example/", words="our website")
        if "titlefont" in done:
            master_title_font(prs, "Georgia")

    def after(p):
        titles = [x["title"] for x in slides]
        icons(p, titles.index("What we offer") + 1, ICON_NAMES)
        if "size" in done:
            size_43_to_169(p)
        if "smartart" in done:
            smartart(p, titles.index("Our process") + 1, ["Reserve", "Paddle", "Return"])
        if "custom" in done:
            titles = [x["title"] for x in slides]
            custom_show(p, "Short version", [titles.index(t) + 1 for t in ("Lakeshore season review", "Highlights", "Rentals", "Visit us")])
        if "personal" in done:
            pres_flags(p, remove_personal=True)
    return build_deck(path, slides, sections_by_title=secs, widescreen=False, title="Lakeshore season review" if "title" in done else "", author=ME,
                      before_save=before, after_save=after)


def build_all(starters):
    out = os.path.join(starters, "week-05")
    os.makedirs(out, exist_ok=True)
    build_l1(os.path.join(out, "P5-L1-review.pptx"))
    build_l2(os.path.join(out, "P5-L2-compare.pptx"))
    build_l2(os.path.join(out, "P5-L2-reviewer.pptx"), {"reviewer"})
    build_l3(os.path.join(out, "P5-L3-show.pptx"))
    build_l4(os.path.join(out, "P5-L4-final.pptx"))
    build_sc5(os.path.join(out, "P5-SC5-team.pptx"))
    build_cap(os.path.join(out, "P5-CAP-exam.pptx"))
