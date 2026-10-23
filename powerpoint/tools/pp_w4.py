# -*- coding: utf-8 -*-
"""
Week 4 practice presentations: Motion, Media and Links.

build(path) makes the starter; build(path, done={...}) makes it with those
steps finished, for testing the checks. Everything is fictional.
"""
import os
from pptkit import *
from pptx.enum.shapes import MSO_SHAPE

ORG = "Lakeshore Paddle Company"
SITE = "https://www.lakeshorepaddle.example/rentals"


def photo_slide(title, n):
    return dict(layout="Title Only", title=title, fn=lambda prs, s: picture(s, media("lake%d.jpg" % n), 2.2, 1.5, 9, alt="Photo: " + title.lower()))


# ---------------------------------------------------------------------------
# Lesson 1: Transitions  (D4.01 fade and 3D; D4.02 timing and sound)
# ---------------------------------------------------------------------------
def build_l1(path, done=frozenset()):
    slides = [
        dict(layout="Title Slide", title="Lakeshore photo tour", subtitle=ORG),
        photo_slide("Morning on the lake", 1),
        photo_slide("The boathouse", 2),
        photo_slide("Sunset paddle", 3),
        dict(layout="Title Only", title="See you on the water"),
    ]

    def before(prs):
        for i, s in enumerate(prs.slides):
            photo = 1 <= i <= 3
            after = 5 if ("timing" in done and photo) else None
            if i == 0 and "3d" in done:
                transition(s, "vortex", dur=2 if "duration" in done else None, option={"dir": "r"} if "option" in done else None)
            elif "fade" in done:
                transition(s, "fade", after=after, sound=("sound" in done and i == 4))
            elif after or ("sound" in done and i == 4):
                transition(s, "cut", after=after, sound=("sound" in done and i == 4))
    return build_deck(path, slides, title="", before_save=before)


# ---------------------------------------------------------------------------
# Lesson 2: Animations  (D4.03 effects; D4.04 timing and order; D4.05 motion path)
# ---------------------------------------------------------------------------
def build_l2(path, done=frozenset()):
    refs = {}

    def why(prs, s):
        refs["list"] = [p for p in s.placeholders if p.placeholder_format.idx == 1][0]
        refs["logo"] = picture(s, media("logo.png"), 10.4, 0.3, 2.5, name="Logo", alt="Lakeshore Paddle Company logo")
        refs["note"] = autoshape(s, MSO_SHAPE.FOLDED_CORNER, 9.2, 4.2, 3.2, 1.6, "Book by Friday!", name="Sticky note")

    def route(prs, s):
        refs["kayak"] = picture(s, media("icon-paddle.png"), 9.5, 3.2, 1.6, 1.6, name="Kayak", alt="Kayak")
        autoshape(s, MSO_SHAPE.RECTANGLE, 0.5, 3.4, 1.6, 1.2, "Dock", name="Dock")

    slides = [
        dict(layout="Title Slide", title="Rent with us", subtitle=ORG),
        dict(title="Why rent with us", body=["Safe gear, checked every day", "Friendly guides", "Boats for every age"], fn=why),
        dict(layout="Title Only", title="Our route", fn=route),
    ]

    def before(prs):
        effects = []
        if "entrance" in done:
            effects.append(dict(shape=refs["list"], cls="entr", preset=10, paragraphs=True, dur=1 if "duration" in done else 0.5))
        if "emphasis" in done:
            e = dict(shape=refs["logo"], cls="emph", preset=8)
            if "after" in done:
                e.update(start="after", delay=1)
            if "reorder" in done:
                effects.insert(0, e)
            else:
                effects.append(e)
        if "exit" in done:
            effects.append(dict(shape=refs["note"], cls="exit", preset=10))
        if effects:
            animate(prs.slides[1], effects)
        if "path" in done:
            animate(prs.slides[2], [dict(shape=refs["kayak"], cls="path", preset=42, path="M 0 0 L -0.6 0 E" if "reverse" in done else "M 0 0 L 0.6 0 E")])
    return build_deck(path, slides, title="", before_save=before)


# ---------------------------------------------------------------------------
# Lesson 3: Video, Audio and 3D  (D4.06 video and music; D4.07 screen recording; D4.08 3D model)
# ---------------------------------------------------------------------------
def build_l3(path, done=frozenset()):
    def tour(prs, s):
        if "video" in done:
            mv = s.shapes.add_movie(media("tour.mp4"), Inches(2.2), Inches(1.5), Inches(8.9), Inches(5), poster_frame_image=media("poster.png"), mime_type="video/mp4")
            mv.name = "tour.mp4"
            s._ot_movie = mv

    def last(prs, s):
        if "recording" in done:
            mv = s.shapes.add_movie(media("tour.mp4"), Inches(2.2), Inches(1.5), Inches(8.9), Inches(5), poster_frame_image=media("poster.png"), mime_type="video/mp4")
            mv.name = "Screen Recording 1"

    slides = [
        dict(layout="Title Slide", title="Welcome to Lakeshore", subtitle=ORG),
        dict(layout="Title Only", title="Tour video", fn=tour),
        dict(layout="Title Only", title="Our new kayak"),
        dict(layout="Title Only", title="How to book online", fn=last),
    ]

    def before(prs):
        s1, s2, s3 = prs.slides[0], prs.slides[1], prs.slides[2]
        if "music" in done:
            a = add_audio(s1, media("music.m4a"))
            media_options(s1, a, start="auto", loop=True, across=True, audio=True)
        if "video" in done:
            mv = [sh for sh in s2.shapes if sh.name == "tour.mp4"][0]
            media_options(s2, mv, trim=(1000, 6000) if "trim" in done else None, fade_in=1000 if "fade" in done else None, start="auto" if "auto" in done else "click")
        if "turntable" in done:
            animate(s3, [dict(shape=Ref(91), cls="emph", preset=8, dur=3)])

    def after(p):
        if "model" in done:
            model3d(p, 3, media("kayak.glb"), rotated="view" in done)
    return build_deck(path, slides, title="", before_save=before, after_save=after)


# ---------------------------------------------------------------------------
# Lesson 4: Links and Zoom  (D4.09 links and action buttons; D4.10 Zoom)
# ---------------------------------------------------------------------------
def _link_text(s, url=None, slide_part=None, words=None):
    """Insert > Link on the words in a slide's text: to a web page, or to a slide in this presentation."""
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if words in r.text:
                    if url:
                        r.hyperlink.address = url
                    else:
                        rid = s.part.relate_to(slide_part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide")
                        rPr = r._r.get_or_add_rPr()
                        rPr.append(X('<a:hlinkClick r:id="%s" action="ppaction://hlinksldjump"/>' % rid)[0])
                    return


def build_l4(path, done=frozenset(), sc=False):
    slides = [
        dict(layout="Title Slide", title="Lakeshore tour", subtitle=ORG),
        dict(title="Menu", body=["Boats", "Prices", "Book online"]),
        dict(title="Kayaks", body=["Single and tandem"]),
        dict(title="Paddleboards", body=["For calm mornings"]),
        dict(title="Prices", body=["Kayak: $20 an hour", "Paddleboard: $18 an hour"]),
        dict(title="Book online", body=["Reserve on our rental page", "Pay at the desk"]),
    ]
    secs = [("Welcome", "Lakeshore tour"), ("Boats", "Kayaks"), ("Prices", "Prices"), ("Book", "Book online")]

    def before(prs):
        S = list(prs.slides)
        if "weblink" in done:
            _link_text(S[5], url=SITE, words="rental page")
        if "slidelink" in done:
            _link_text(S[1], slide_part=S[4].part, words="Prices")
        if "button" in done:
            b = autoshape(S[4], MSO_SHAPE.ACTION_BUTTON_HOME, 11.6, 6, 1, 0.8, name="Action Button: Go Home 3")
            rid = S[4].part.relate_to(S[1].part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide")
            nv = b._element.find(".//" + qn("p:cNvPr"))
            nv.append(X('<a:hlinkClick r:id="%s" action="ppaction://hlinksldjump" highlightClick="1"/>' % rid)[0])

    def after(p):
        if "summary" in done:
            zoom(p, 2, "summary", [1, 2, 3], return_to=True, shape_id=90)
        if "section" in done:
            zoom(p, 2, "section", [2], return_to=True, shape_id=94)
        if "slidezoom" in done:
            zoom(p, 3, "slide", [4], return_to="return" in done, shape_id=95)
    return build_deck(path, slides, sections_by_title=secs, title="", before_save=before, after_save=after)


# ---------------------------------------------------------------------------
# Skill Check 4: kiosk tour
# ---------------------------------------------------------------------------
def build_sc4(path, done=frozenset()):
    refs = {}

    def boats(prs, s):
        refs["list"] = [p for p in s.placeholders if p.placeholder_format.idx == 1][0]
        refs["logo"] = picture(s, media("logo.png"), 10.4, 0.3, 2.5, name="Logo", alt="Lakeshore Paddle Company logo")

    def route(prs, s):
        refs["kayak"] = picture(s, media("icon-paddle.png"), 1, 3.2, 1.6, 1.6, name="Kayak", alt="Kayak")

    def tour(prs, s):
        if "video" in done:
            mv = s.shapes.add_movie(media("tour.mp4"), Inches(2.2), Inches(1.5), Inches(8.9), Inches(5), poster_frame_image=media("poster.png"), mime_type="video/mp4")
            mv.name = "tour.mp4"

    slides = [
        dict(layout="Title Slide", title="Lakeshore kiosk tour", subtitle=ORG),
        dict(title="Menu", body=["Our boats", "The route", "Tour", "Book"]),
        dict(title="Our boats", body=["Kayaks", "Canoes", "Paddleboards"], fn=boats),
        dict(layout="Title Only", title="The route", fn=route),
        dict(layout="Title Only", title="Tour", fn=tour),
        dict(title="Book", body=["Book now on our rental page"]),
    ]
    secs = [("Start", "Lakeshore kiosk tour"), ("Boats", "Our boats"), ("Tour", "The route"), ("Book", "Book")]

    def before(prs):
        S = list(prs.slides)
        for i, s in enumerate(S):
            if "push" in done or "after" in done:
                transition(s, "push" if "push" in done else "cut", option={"dir": "r"} if "pushoption" in done else None, after=8 if "after" in done else None)
        eff = []
        if "flyin" in done:
            eff.append(dict(shape=refs["list"], cls="entr", preset=2, subtype=4, paragraphs=True, start="after" if "listafter" in done else "click",
                            next_start="after" if "listafter" in done else "click"))
        if "spin" in done:
            e = dict(shape=refs["logo"], cls="emph", preset=8)
            eff.insert(0, e) if "first" in done else eff.append(e)
        if eff:
            animate(S[2], eff)
        if "path" in done:
            animate(S[3], [dict(shape=refs["kayak"], cls="path", preset=42, path="M 0 0 L 0.6 0 E")])
        if "video" in done:
            mv = [sh for sh in S[4].shapes if sh.name == "tour.mp4"][0]
            media_options(S[4], mv, trim=(1000, 6000) if "trim" in done else None, start="auto" if "auto" in done else "click")
        if "weblink" in done:
            _link_text(S[5], url=SITE, words="Book now")

    def after(p):
        if "summary" in done:
            zoom(p, 2, "summary", [1, 2, 3], return_to=True, shape_id=90)
        if "slidezoom" in done:
            zoom(p, 2, "slide", [6], return_to=True, shape_id=95)
    return build_deck(path, slides, sections_by_title=secs, title="", before_save=before, after_save=after)


def build_all(starters):
    import shutil
    out = os.path.join(starters, "week-04")
    os.makedirs(out, exist_ok=True)
    for name in ("tour.mp4", "music.m4a", "kayak.glb"):
        shutil.copy(media(name), os.path.join(out, name))
    build_l1(os.path.join(out, "P4-L1-transitions.pptx"))
    build_l2(os.path.join(out, "P4-L2-animations.pptx"))
    build_l3(os.path.join(out, "P4-L3-media.pptx"))
    build_l4(os.path.join(out, "P4-L4-links.pptx"))
    build_sc4(os.path.join(out, "P4-SC4-kiosk.pptx"))
