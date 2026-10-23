# -*- coding: utf-8 -*-
"""
Week 3 practice presentations: Visual Slides.

build(path) makes the starter; build(path, done={...}) makes it with those
steps finished, for testing the checks. Everything is fictional.
"""
import os
from pptkit import *
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE

ORG = "Lakeshore Paddle Company"
ICON_NAMES = ["Icon wave", "Icon sun", "Icon paddle"]


def _icons_row(s, done, top=2.4):
    """Three icons (wave, sun, paddle) at uneven heights and gaps, as Insert > Icons leaves them."""
    spots = [(1.2, top), (4.1, top + 0.45), (8.3, top - 0.3)] if "align" not in done else [(1.2, top), (5.2, top), (9.2, top)]
    if "align" in done and "distribute" not in done:
        spots = [(1.2, top), (4.1, top), (8.3, top)]
    pics = []
    for (x, y), kind in zip(spots, ["wave", "sun", "paddle"]):
        pics.append(picture(s, media("icon-%s.png" % kind), x, y, 1.4, 1.4, name="Icon %s" % kind, decorative=True))
    if "group" in done:
        group(s, pics, name="Icons")
    return pics


# ---------------------------------------------------------------------------
# Lesson 1: Pictures  (D3.01 insert, resize, crop; D3.02 styles, alt text, screen clipping)
# ---------------------------------------------------------------------------
def build_l1(path, done=frozenset()):
    def fleet(prs, s):
        if "insert" in done:
            pic = picture(s, media("lake3.jpg"), 1, 1.6, 6 if "resize" in done else None)
            if "stretch" in done:
                pic.width = Inches(7); pic.height = Inches(3)
            if "crop" in done:
                crop(pic, 1 / 6, 0, 1 / 6, 0)          # 900x600 -> 600x600 shown
                pic.width = pic.height                  # PowerPoint shrinks the frame to match
            if "style" in done:
                picture_style(pic)
            if "alt" in done:
                set_alt(pic, "A red kayak on a calm lake below green hills")

    def beach(prs, s):
        pic = picture(s, media("lake1.jpg"), 1, 1.6, 5.5, name="Beach photo")
        if "alt" in done:
            set_alt(pic, "The Lakeshore beach with boats lined up on the sand")

    def weather(prs, s):
        if "clipping" in done:
            picture(s, media("lake4.jpg"), 2, 2, 4, name="Picture 3")

    slides = [
        dict(layout="Title Slide", title="Our boats and beach", subtitle=ORG),
        dict(layout="Title Only", title="Our fleet", fn=fleet),
        dict(layout="Title Only", title="The beach", fn=beach),
        dict(layout="Title Only", title="Today's weather", fn=weather),
    ]
    return build_deck(path, slides, title="")


# ---------------------------------------------------------------------------
# Lesson 2: Shapes, Order and Alignment  (D3.03 shapes and text boxes; D3.04 order, align, group, guides, ink)
# ---------------------------------------------------------------------------
def build_l2(path, done=frozenset()):
    def find_us(prs, s):
        if "callout" in done:
            w, h = (3, 1) if "size" in done else (2.2, 1.3)
            c = autoshape(s, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, 7, 2, w, h, "Free demos Saturday")
            if "shapestyle" in done:
                shape_style(c)
        if "textbox" in done:
            text_box(s, 1, 5.5, 4, 0.8, ["Open 9 to 7 every day"], size=24)

    def rentals(prs, s):
        logo = picture(s, media("logo.png"), 0.6, 5.6, 2.4, name="Logo", alt="Lakeshore Paddle Company logo")
        banner = autoshape(s, MSO_SHAPE.RECTANGLE, 0, 5.4, 13.333, 1.2, "Rent by the hour or the day", name="Banner")
        if "front" in done:
            to_front(s, logo)
        _icons_row(s, done)

    slides = [
        dict(layout="Title Slide", title="Find us on the beach", subtitle=ORG),
        dict(layout="Title Only", title="Find us", fn=find_us),
        dict(layout="Title Only", title="Rentals", fn=rentals),
    ]
    def after(p):
        icons(p, 3, ICON_NAMES)
        if "guides" in done:
            last_view(p, "sldThumbnailView", guides=True)
        if "ink" in done:
            ink(p, 3)
    return build_deck(path, slides, title="", after_save=after)


# ---------------------------------------------------------------------------
# Lesson 3: Tables and Charts  (D3.05 table; D3.06 chart)
# ---------------------------------------------------------------------------
RENTALS = [("Kayak", (120, 180, 150)), ("Canoe", (60, 95, 80)), ("Paddleboard", (90, 140, 110))]


def add_chart(s, kind="column", title=None, labels=False, alt=None):
    cd = CategoryChartData()
    cd.categories = ["June", "July", "August"]
    for name, vals in RENTALS:
        cd.add_series(name, vals)
    t = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}[kind]
    gf = s.shapes.add_chart(t, Inches(1.5), Inches(1.6), Inches(9), Inches(5), cd)
    ch = gf.chart
    ch.has_legend = True
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
    else:
        ch.has_title = False
    if labels:
        ch.plots[0].has_data_labels = True
    if alt:
        set_alt(gf, alt)
    return gf


def build_l3(path, done=frozenset()):
    def prices(prs, s):
        if "table" in done:
            rows = [["Boat", "1 hour", "Half day", "Full day"], ["Kayak", "$20", "$45", "$70"], ["Canoe", "$25", "$55", "$85"]]
            if "addrow" in done:
                rows.append(["Paddleboard", "$18", "$40", "$60"])
            if "delcol" in done:
                rows = [r[:3] for r in rows]
            t = s.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(1.8), Inches(9), Inches(0.5 * len(rows)))
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.table.cell(i, j).text = v
            if "tablestyle" in done:
                table_style(t, MEDIUM_2_ACCENT_2)

    def season(prs, s):
        if "chart" in done:
            add_chart(s, "bar" if "charttype" in done else "column", "Rentals by month" if "charttitle" in done else None,
                      labels="charttitle" in done, alt="Kayaks were rented most every month, with July the busiest." if "chartalt" in done else None)

    slides = [
        dict(layout="Title Slide", title="Prices and rentals", subtitle=ORG),
        dict(layout="Title Only", title="Prices", fn=prices),
        dict(layout="Title Only", title="Season rentals", fn=season,
             notes="Rentals: Kayak 120, 180, 150; Canoe 60, 95, 80; Paddleboard 90, 140, 110 (June, July, August)."),
    ]
    return build_deck(path, slides, title="")


# ---------------------------------------------------------------------------
# Lesson 4: SmartArt  (D3.07)
# ---------------------------------------------------------------------------
STEPS = ["Check the weather", "Put on a life jacket", "Tell someone your plan"]


def build_l4(path, done=frozenset()):
    slides = [
        dict(layout="Title Slide", title="Paddle safely", subtitle=ORG),
        dict(title="Safety steps", body=None if "smartart" in done else STEPS),
        dict(title="Our values", body=["Safety", "Fun", "Respect for the lake"]),
    ]

    def after(p):
        if "smartart" in done:
            items = STEPS + (["Have fun"] if "addstep" in done else [])
            smartart(p, 2, items, layout_id="cycle2" if "layout" in done else "process1", colors="colorful1" if "colors" in done else "accent1_2",
                     alt="Four safety steps in a cycle" if "smartalt" in done else "")
    return build_deck(path, slides, title="", after_save=after)


# ---------------------------------------------------------------------------
# Skill Check 3: season report
# ---------------------------------------------------------------------------
def build_sc3(path, done=frozenset()):
    def beach(prs, s):
        pic = picture(s, media("lake2.jpg"), 1, 1.6, 6, name="Beach photo")
        if "crop" in done:
            crop(pic, 1 / 6, 0, 1 / 6, 0); pic.width = pic.height
        if "style" in done:
            picture_style(pic)
        if "alt" in done:
            set_alt(pic, "Families launching kayaks from the beach")

    def highlights(prs, s):
        logo = picture(s, media("logo.png"), 9.8, 0.3, 2.6, name="Logo", alt="Lakeshore Paddle Company logo")
        box = autoshape(s, MSO_SHAPE.RECTANGLE, 9.2, 0.2, 4, 1.2, name="Corner box")
        if "front" in done:
            to_front(s, logo)
        _icons_row(s, done, top=4.8)
        if "shape" in done:
            a = autoshape(s, MSO_SHAPE.RIGHT_ARROW, 1, 2, 4 if "size" in done else 3, 1.4, "Busiest day: July 4")
            if "shapestyle" in done:
                shape_style(a)

    def prices(prs, s):
        if "table" in done:
            rows = [["Boat", "Hour", "Day"], ["Kayak", "$20", "$70"], ["Canoe", "$25", "$85"]]
            t = s.shapes.add_table(3, 3, Inches(1), Inches(1.8), Inches(8), Inches(1.5))
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.table.cell(i, j).text = v
            if "tablestyle" in done:
                table_style(t, LIGHT_1_ACCENT_1)

    def rentals(prs, s):
        if "chart" in done:
            add_chart(s, "column", "Rentals by month" if "charttitle" in done else None)

    slides = [
        dict(layout="Title Slide", title="Season report", subtitle=ORG),
        dict(layout="Title Only", title="Our beach", fn=beach),
        dict(layout="Title Only", title="Highlights", fn=highlights),
        dict(layout="Title Only", title="Prices", fn=prices),
        dict(layout="Title Only", title="Rentals", fn=rentals,
             notes="Kayak 120, 180, 150; Canoe 60, 95, 80; Paddleboard 90, 140, 110 (June, July, August)."),
        dict(title="Steps", body=None if "smartart" in done else ["Reserve online", "Check in at the desk", "Get your gear"]),
    ]

    def after(p):
        icons(p, 3, ICON_NAMES)
        if "smartart" in done:
            smartart(p, 6, ["Reserve online", "Check in at the desk", "Get your gear", "Paddle!"] if "addshape" in done else ["Reserve online", "Check in at the desk", "Get your gear"])
    return build_deck(path, slides, title="", after_save=after)


def build_all(starters):
    out = os.path.join(starters, "week-03")
    os.makedirs(out, exist_ok=True)
    make_photo(os.path.join(out, "kayak.jpg"), seed=3)
    build_l1(os.path.join(out, "P3-L1-pictures.pptx"))
    build_l2(os.path.join(out, "P3-L2-shapes.pptx"))
    build_l3(os.path.join(out, "P3-L3-data.pptx"))
    build_l4(os.path.join(out, "P3-L4-smartart.pptx"))
    build_sc3(os.path.join(out, "P3-SC3-report.pptx"))
