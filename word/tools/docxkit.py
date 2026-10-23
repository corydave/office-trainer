# docxkit.py - small helpers for building Word files that python-docx can't make by itself.
#
# make-starters.py (the practice documents) and make-test-docs.py (the answer keys)
# both use these. python-docx handles paragraphs, runs, tables, pictures and
# sections; everything else (fields, bookmarks, links, footnotes, comments,
# tracked changes, lists with special numbering, shapes, SmartArt, charts...)
# is written here as the raw Word XML the checker reads.
#
# Everything in these files is fictional: businesses, people, data and sources.
import copy
import datetime
import itertools
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.opc.packuri import PackURI

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
NSW = 'xmlns:w="%s" xmlns:w14="%s" xmlns:w15="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"' % (W, W14, W15)
DATE = "2026-10-20T09:00:00Z"
_ids = itertools.count(100)   # unique ids for revisions, comments and bookmarks


def next_id():
    return str(next(_ids))


def esc(text):
    """Make text safe inside XML."""
    return (str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


# --------------------------------------------------------------------------- documents
def new_doc(font="Aptos", size=11, after=8, lang="en-US", compat=15):
    """A blank document: known Normal style, Letter size, 1 in margins, empty properties.
    compat=15 is current Word. (python-docx's own template is Word 2010 mode, 14,
    which makes Word show "Compatibility Mode" in the title bar.)"""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), font)
    set_lang(rpr, lang)
    pf = normal.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(after); pf.line_spacing = 1.0
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(1))
    cp = doc.core_properties
    cp.author = ""; cp.title = ""; cp.keywords = ""; cp.subject = ""; cp.comments = ""; cp.last_modified_by = ""
    set_compat(doc, compat)
    return doc


def set_compat(doc, mode):
    for c in doc.settings.element.iter(qn("w:compatSetting")):
        if c.get(qn("w:name")) == "compatibilityMode":
            c.set(qn("w:val"), str(mode))


def set_lang(rpr, lang):
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang"); rpr.append(el)
    el.set(qn("w:val"), lang)


def setting(doc, xml):
    """Add one element to settings.xml, in the order Word expects (good enough: before w:compat)."""
    el = parse_xml(xml)
    compat = doc.settings.element.find(qn("w:compat"))
    if compat is not None:
        compat.addprevious(el)
    else:
        doc.settings.element.append(el)
    return el


def para(doc, text="", style=None, bold=False, italic=False, size=None, color=None, font=None,
         align=None, after=None, before=None, keep_next=False, container=None):
    """Add a paragraph (to the document, or to a header/footer/cell given as container)."""
    target = container if container is not None else doc
    p = target.add_paragraph(style=style) if style else target.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold or None
        r.italic = italic or None
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if font: r.font.name = font
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    if keep_next: p.paragraph_format.keep_with_next = True
    return p


def find(doc, starts, paragraphs=None):
    for p in paragraphs if paragraphs is not None else doc.paragraphs:
        if p.text.startswith(starts):
            return p
    raise ValueError("no paragraph starts with " + starts)


def raw_run(p, xml):
    """Append raw run-level XML (w:r, w:hyperlink, w:fldSimple...) to a paragraph."""
    el = parse_xml(xml)
    p._p.append(el)
    return el


def add_part(doc, partname, content_type, xml, reltype):
    """Add a new part (footnotes.xml, comments.xml...) and link it to the document. Returns the rId."""
    part = Part(PackURI(partname), content_type, xml.encode("utf-8") if isinstance(xml, str) else xml, doc.part.package)
    return doc.part.relate_to(part, reltype)


# --------------------------------------------------------------------------- fields and links
def field_xml(code, result="", rpr=""):
    """A complex field as runs: begin, instruction, separate, the shown result, end."""
    out = '<w:r %s>%s<w:fldChar w:fldCharType="begin"/></w:r>' % (NSW, rpr)
    out += '<w:r %s>%s<w:instrText xml:space="preserve"> %s </w:instrText></w:r>' % (NSW, rpr, esc(code))
    out += '<w:r %s>%s<w:fldChar w:fldCharType="separate"/></w:r>' % (NSW, rpr)
    if result:
        out += '<w:r %s>%s<w:t xml:space="preserve">%s</w:t></w:r>' % (NSW, rpr, esc(result))
    out += '<w:r %s>%s<w:fldChar w:fldCharType="end"/></w:r>' % (NSW, rpr)
    return out


def add_field(p, code, result=""):
    for part in field_xml(code, result).split("</w:r>")[:-1]:
        raw_run(p, part + "</w:r>")


def add_simple_field(p, code, result=""):
    raw_run(p, '<w:fldSimple %s w:instr=" %s "><w:r><w:t>%s</w:t></w:r></w:fldSimple>' % (NSW, esc(code), esc(result)))


def bookmark(p, name):
    """Put a bookmark around a whole paragraph."""
    bid = next_id()
    p._p.insert(1 if p._p.pPr is not None else 0, parse_xml('<w:bookmarkStart %s w:id="%s" w:name="%s"/>' % (NSW, bid, esc(name))))
    p._p.append(parse_xml('<w:bookmarkEnd %s w:id="%s"/>' % (NSW, bid)))


def hyperlink(doc, p, text, url=None, anchor=None):
    """Add a hyperlink run: to a web address (url) or to a bookmark (anchor)."""
    if url:
        rid = doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
        attrs = 'r:id="%s"' % rid
    else:
        attrs = 'w:anchor="%s"' % esc(anchor)
    raw_run(p, '<w:hyperlink %s %s w:history="1"><w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr><w:t xml:space="preserve">%s</w:t></w:r></w:hyperlink>' % (NSW, attrs, esc(text)))


def page_break_before(p):
    p.paragraph_format.page_break_before = True


def column_break(p):
    raw_run(p, '<w:r %s><w:br w:type="column"/></w:r>' % NSW)


# --------------------------------------------------------------------------- sections
def set_columns(section, n, space=0.5):
    sp = section._sectPr
    cols = sp.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sp.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space * 1440)))


def set_section_type(section, kind):
    """kind: nextPage, continuous, evenPage, oddPage."""
    sp = section._sectPr
    t = sp.find(qn("w:type"))
    if t is None:
        t = OxmlElement("w:type"); sp.insert(0, t)
    t.set(qn("w:val"), kind)


def page_border(section):
    section._sectPr.append(parse_xml(
        '<w:pgBorders %s w:offsetFrom="page"><w:top w:val="single" w:sz="12" w:space="24" w:color="0E7C86"/>'
        '<w:left w:val="single" w:sz="12" w:space="24" w:color="0E7C86"/><w:bottom w:val="single" w:sz="12" w:space="24" w:color="0E7C86"/>'
        '<w:right w:val="single" w:sz="12" w:space="24" w:color="0E7C86"/></w:pgBorders>' % NSW))


def page_color(doc, hex6):
    doc.element.insert(0, parse_xml('<w:background %s w:color="%s"/>' % (NSW, hex6)))
    setting(doc, '<w:displayBackgroundShape %s/>' % NSW)


def watermark(section, text="DRAFT"):
    """A text watermark in the header, drawn the way Word does it (VML WordArt)."""
    hp = section.header.paragraphs[0]
    raw_run(hp, ('<w:r %s xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><w:pict>'
                 '<v:shape id="PowerPlusWaterMarkObject1" o:spid="_x0000_s2049" type="#_x0000_t136" style="position:absolute;width:468pt;height:117pt;rotation:315;z-index:-251657216" fillcolor="silver" stroked="f">'
                 '<v:textpath style="font-family:&quot;Calibri&quot;;font-size:1pt" string="%s"/></v:shape></w:pict></w:r>') % (NSW, esc(text)))


def drop_cap(p, lines=3):
    p._p.get_or_add_pPr().insert(0, parse_xml('<w:framePr %s w:dropCap="drop" w:lines="%d" w:wrap="around" w:vAnchor="text" w:hAnchor="text"/>' % (NSW, lines)))


# --------------------------------------------------------------------------- lists
def add_numbering(doc, fmt="decimal", text="%1.", start=1, bullet_font=None, levels=None):
    """Add a list design (abstractNum) and one list (num) that uses it. Returns the numId.
    levels: optional list of (fmt, text) for levels 2+; defaults are a, b, c / i, ii, iii."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    aid = max(abstract_ids + [0]) + 1
    nid = max(num_ids + [0]) + 1
    lvls = [(fmt, text)] + (levels or [("lowerLetter", "%2."), ("lowerRoman", "%3.")])
    xml = '<w:abstractNum %s w:abstractNumId="%d"><w:multiLevelType w:val="hybridMultilevel"/>' % (NSW, aid)
    for i, (f, t) in enumerate(lvls):
        font = '<w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s" w:hint="default"/></w:rPr>' % (bullet_font, bullet_font) if (f == "bullet" and bullet_font) else ""
        xml += ('<w:lvl w:ilvl="%d"><w:start w:val="%d"/><w:numFmt w:val="%s"/><w:lvlText w:val="%s"/><w:lvlJc w:val="left"/>'
                '<w:pPr><w:ind w:left="%d" w:hanging="360"/></w:pPr>%s</w:lvl>') % (i, start if i == 0 else 1, f, esc(t), 720 * (i + 1), font)
    xml += "</w:abstractNum>"
    first_num = numbering.find(qn("w:num"))
    el = parse_xml(xml)
    if first_num is not None:
        first_num.addprevious(el)
    else:
        numbering.append(el)
    numbering.append(parse_xml('<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>' % (NSW, nid, aid)))
    return nid


def restart_num(doc, base_num_id, start):
    """A new list instance of the same design that starts again at `start` (Restart / Set Numbering Value)."""
    numbering = doc.part.numbering_part.element
    base = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(base_num_id)][0]
    aid = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    nid = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    numbering.append(parse_xml('<w:num %s w:numId="%d"><w:abstractNumId w:val="%s"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="%d"/></w:lvlOverride></w:num>' % (NSW, nid, aid, start)))
    return nid


def set_list(p, num_id, level=0):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn("w:numPr"))
    if old is not None:
        pPr.remove(old)
    pPr.append(parse_xml('<w:numPr %s><w:ilvl w:val="%d"/><w:numId w:val="%d"/></w:numPr>' % (NSW, level, num_id)))


# --------------------------------------------------------------------------- notes
def footnotes(doc, notes, fmt=None):
    """notes: list of (paragraph, note text). Adds footnotes.xml and the reference marks."""
    body = ('<w:footnotes %s><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>') % NSW
    for i, (p, text) in enumerate(notes, start=1):
        body += ('<w:footnote w:id="%d"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
                 '<w:r><w:t xml:space="preserve"> %s</w:t></w:r></w:p></w:footnote>') % (i, esc(text))
        raw_run(p, '<w:r %s><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="%d"/></w:r>' % (NSW, i))
    body += "</w:footnotes>"
    add_part(doc, "/word/footnotes.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", body,
             "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes")
    if fmt:
        doc.sections[0]._sectPr.append(parse_xml('<w:footnotePr %s><w:numFmt w:val="%s"/></w:footnotePr>' % (NSW, fmt)))


# --------------------------------------------------------------------------- comments
def comments(doc, items):
    """items: list of dicts {anchor: paragraph, author, text, reply_to: index, done: bool}.
    Writes comments.xml and commentsExtended.xml (replies and resolved state)."""
    xml = '<w:comments %s>' % NSW
    ext = '<w15:commentsEx %s>' % NSW
    para_ids = []
    for i, c in enumerate(items):
        cid = str(i)
        pid = "%08X" % (0x10000000 + i)
        para_ids.append(pid)
        initials = "".join(w[0] for w in c["author"].split())
        xml += ('<w:comment w:id="%s" w:author="%s" w:date="%s" w:initials="%s"><w:p w14:paraId="%s" w14:textId="77777777">'
                '<w:r><w:annotationRef/></w:r><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p></w:comment>') % (cid, esc(c["author"]), DATE, initials, pid, esc(c["text"]))
        parent = ' w15:paraIdParent="%s"' % para_ids[c["reply_to"]] if c.get("reply_to") is not None else ""
        ext += '<w15:commentEx w15:paraId="%s"%s w15:done="%s"/>' % (pid, parent, "1" if c.get("done") else "0")
        p = c["anchor"]._p
        # The range covers the whole paragraph; a reply shares its parent's range.
        start = parse_xml('<w:commentRangeStart %s w:id="%s"/>' % (NSW, cid))
        pPr = p.find(qn("w:pPr"))
        if pPr is not None:
            pPr.addnext(start)
        else:
            p.insert(0, start)
        p.append(parse_xml('<w:commentRangeEnd %s w:id="%s"/>' % (NSW, cid)))
        p.append(parse_xml('<w:r %s><w:commentReference w:id="%s"/></w:r>' % (NSW, cid)))
    xml += "</w:comments>"
    ext += "</w15:commentsEx>"
    add_part(doc, "/word/comments.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml", xml,
             "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments")
    add_part(doc, "/word/commentsExtended.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml", ext,
             "http://schemas.microsoft.com/office/2011/relationships/commentsExtended")


# --------------------------------------------------------------------------- tracked changes
def tracked_insert(p, text, author):
    raw_run(p, '<w:ins %s w:id="%s" w:author="%s" w:date="%s"><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:ins>' % (NSW, next_id(), esc(author), DATE, esc(text)))


def tracked_delete(p, text, author):
    raw_run(p, '<w:del %s w:id="%s" w:author="%s" w:date="%s"><w:r><w:delText xml:space="preserve">%s</w:delText></w:r></w:del>' % (NSW, next_id(), esc(author), DATE, esc(text)))


def tracked_bold(run, author):
    """Make a run bold as a tracked formatting change."""
    rpr = run._r.get_or_add_rPr()
    rpr.append(parse_xml('<w:b %s/>' % NSW))
    rpr.append(parse_xml('<w:rPrChange %s w:id="%s" w:author="%s" w:date="%s"><w:rPr/></w:rPrChange>' % (NSW, next_id(), esc(author), DATE)))


def tracking_on(doc, locked=False):
    setting(doc, '<w:trackRevisions %s/>' % NSW)
    if locked:
        protect(doc, "trackedChanges", password=True)


def protect(doc, edit, password=False):
    pw = ' w:cryptProviderType="rsaAES" w:cryptAlgorithmClass="hash" w:cryptAlgorithmType="typeAny" w:cryptAlgorithmSid="14" w:cryptSpinCount="100000" w:hash="dGVzdA==" w:salt="c2FsdA=="' if password else ""
    setting(doc, '<w:documentProtection %s w:edit="%s" w:enforcement="1"%s/>' % (NSW, edit, pw))


# --------------------------------------------------------------------------- tables
def header_row(table):
    """Mark the first row as a header row that repeats on each page."""
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml('<w:tblHeader %s/>' % NSW))


def cell_margins(table, inches):
    tblPr = table._tbl.tblPr
    w = int(inches * 1440)
    tblPr.append(parse_xml('<w:tblCellMar %s><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tblCellMar>' % (NSW, w, w)))


def column_width(table, col, inches):
    tw = int(inches * 1440)
    table._tbl.tblGrid.findall(qn("w:gridCol"))[col].set(qn("w:w"), str(tw))
    for row in table.rows:
        cell = row.cells[col]
        cell.width = Inches(inches)


# --------------------------------------------------------------------------- drawings
def _anchor_xml(kind_xml, name, alt="", wrap="square", halign="right", valign="top", width=2.0, height=1.2, decorative=False, did=None):
    did = did or next_id()
    dec = ('<a:extLst><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}"><adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/></a:ext></a:extLst>' if decorative else "")
    wraps = {"square": '<wp:wrapSquare wrapText="bothSides"/>', "tight": '<wp:wrapTight wrapText="bothSides"><wp:wrapPolygon edited="0"><wp:start x="0" y="0"/><wp:lineTo x="0" y="21600"/></wp:wrapPolygon></wp:wrapTight>',
             "topAndBottom": "<wp:wrapTopAndBottom/>", "none": "<wp:wrapNone/>"}
    cx, cy = int(width * 914400), int(height * 914400)
    return ('<w:r %s xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><w:drawing>'
            '<wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
            '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>%s</wp:align></wp:positionH>'
            '<wp:positionV relativeFrom="margin"><wp:align>%s</wp:align></wp:positionV><wp:extent cx="%d" cy="%d"/><wp:effectExtent l="0" t="0" r="0" b="0"/>%s'
            '<wp:docPr id="%s" name="%s" descr="%s">%s</wp:docPr><wp:cNvGraphicFramePr/>'
            '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">%s</a:graphicData></a:graphic>'
            '</wp:anchor></w:drawing></w:r>') % (NSW, halign, valign, cx, cy, wraps[wrap], did, esc(name), esc(alt), dec, kind_xml)


def shape_xml(text="", textbox=False, fill=None, outline=None, geom="rect"):
    sp = '<wps:cNvSpPr txBox="1"/>' if textbox else "<wps:cNvSpPr/>"
    fill_xml = '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill if fill else ""
    line_xml = '<a:ln w="19050"><a:solidFill><a:srgbClr val="%s"/></a:solidFill></a:ln>' % outline if outline else ""
    txbx = ('<wps:txbx><w:txbxContent><w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p></w:txbxContent></wps:txbx>' % esc(text)) if text else ""
    return ('<wps:wsp>%s<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="914400"/></a:xfrm><a:prstGeom prst="%s"><a:avLst/></a:prstGeom>%s%s</wps:spPr>'
            '%s<wps:bodyPr/></wps:wsp>') % (sp, geom, fill_xml, line_xml, txbx)


def add_shape(p, name, text="", textbox=False, fill=None, outline=None, alt="", wrap="square", halign="right", valign="top", decorative=False):
    raw_run(p, _anchor_xml(shape_xml(text, textbox, fill, outline), name, alt=alt, wrap=wrap, halign=halign, valign=valign, decorative=decorative))


def picture_docpr(doc, index):
    """The docPr (name, alt text) of the nth picture python-docx inserted."""
    return list(doc.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"))[index]


def set_alt(doc, index, alt="", decorative=False):
    dp = picture_docpr(doc, index)
    dp.set("descr", alt)
    if decorative:
        dp.append(parse_xml('<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
                            '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/></a:ext></a:extLst>'))


def picture_effects(doc, index, style=False, artistic=False, bg_removed=False):
    """Add a picture style (shadow + frame), an artistic effect, or background removal to the nth picture."""
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    pics = list(doc.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"))
    pic = pics[index]
    blip = next(pic.iter("{%s}blip" % A))
    if artistic or bg_removed:
        effect = ('<a14:artisticPencilSketch/>' if artistic else "") + ('<a14:backgroundRemoval t="10000" b="90000" l="10000" r="90000"/>' if bg_removed else "")
        blip.append(parse_xml('<a:extLst xmlns:a="%s"><a:ext uri="{BEBA8EAE-BF5A-486C-A8C5-ECC9F3942E4B}"><a14:imgProps xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main">'
                              '<a14:imgLayer r:embed="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><a14:imgEffect>%s</a14:imgEffect></a14:imgLayer></a14:imgProps></a:ext></a:extLst>'
                              % (A, blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"), effect)))
    if style:
        spPr = pic.find("{http://schemas.openxmlformats.org/drawingml/2006/picture}spPr")
        spPr.append(parse_xml('<a:ln xmlns:a="%s" w="88900"><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></a:ln>' % A))
        spPr.append(parse_xml('<a:effectLst xmlns:a="%s"><a:outerShdw blurRad="254000" algn="tl" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="43000"/></a:srgbClr></a:outerShdw></a:effectLst>' % A))


def float_picture(doc, index, wrap="square", halign="right", valign="top"):
    """Turn the nth inline picture into a floating one (as Layout Options would)."""
    WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    inline = list(doc.element.body.iter("{%s}inline" % WP))[index]
    extent = inline.find("{%s}extent" % WP)
    docpr = inline.find("{%s}docPr" % WP)
    graphic = inline.find("{http://schemas.openxmlformats.org/drawingml/2006/main}graphic")
    wraps = {"square": '<wp:wrapSquare wrapText="bothSides"/>', "topAndBottom": "<wp:wrapTopAndBottom/>", "none": "<wp:wrapNone/>"}
    anchor = parse_xml(('<wp:anchor xmlns:wp="%s" distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251660288" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
                        '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>%s</wp:align></wp:positionH><wp:positionV relativeFrom="margin"><wp:align>%s</wp:align></wp:positionV>'
                        '<wp:extent cx="%s" cy="%s"/><wp:effectExtent l="0" t="0" r="0" b="0"/>%s<wp:cNvGraphicFramePr/></wp:anchor>') % (WP, halign, valign, extent.get("cx"), extent.get("cy"), wraps[wrap]))
    anchor.insert(len(anchor) - 1, copy.deepcopy(docpr))
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline, anchor)


def add_graphic_inline(p, uri, inner_xml, name, alt="", did=None):
    """An inline drawing of any kind (SmartArt, chart, icon...) given its graphicData uri and inner XML."""
    raw_run(p, ('<w:r %s xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><w:drawing>'
                '<wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="5486400" cy="3200400"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
                '<wp:docPr id="%s" name="%s" descr="%s"/><wp:cNvGraphicFramePr/><a:graphic><a:graphicData uri="%s">%s</a:graphicData></a:graphic></wp:inline></w:drawing></w:r>')
            % (NSW, did or next_id(), esc(name), esc(alt), uri, inner_xml))


def add_smartart(doc, p, items, layout="process1", colors="colorful1", style="simple3", alt=""):
    """A SmartArt graphic: data (the text), layout, colors and style parts, as Word links them."""
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagram"
    n = next_id()
    D = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    pts = '<dgm:pt modelId="{00000000-0000-0000-0000-000000000000}" type="doc"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:endParaRPr/></a:p></dgm:t></dgm:pt>'
    for i, t in enumerate(items):
        pts += '<dgm:pt modelId="{00000000-0000-0000-0000-0000000000%02d}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></dgm:t></dgm:pt>' % (i + 1, esc(t))
    data = '<dgm:dataModel xmlns:dgm="%s" xmlns:a="%s"><dgm:ptLst>%s</dgm:ptLst><dgm:cxnLst/><dgm:bg/><dgm:whole/></dgm:dataModel>' % (D, A, pts)
    lo = '<dgm:layoutDef xmlns:dgm="%s" uniqueId="urn:microsoft.com/office/officeart/2005/8/layout/%s"/>' % (D, layout)
    cs = '<dgm:colorsDef xmlns:dgm="%s" uniqueId="urn:microsoft.com/office/officeart/2005/8/colors/%s"/>' % (D, colors)
    qs = '<dgm:styleDef xmlns:dgm="%s" uniqueId="urn:microsoft.com/office/officeart/2005/8/quickstyle/%s"/>' % (D, style)
    ct = "application/vnd.openxmlformats-officedocument.drawingml.diagram"
    r_dm = add_part(doc, "/word/diagrams/data%s.xml" % n, ct + "Data+xml", data, rel + "Data")
    r_lo = add_part(doc, "/word/diagrams/layout%s.xml" % n, ct + "Layout+xml", lo, rel + "Layout")
    r_qs = add_part(doc, "/word/diagrams/quickStyle%s.xml" % n, ct + "Style+xml", qs, rel + "QuickStyle")
    r_cs = add_part(doc, "/word/diagrams/colors%s.xml" % n, ct + "Colors+xml", cs, rel + "Colors")
    add_graphic_inline(p, D, '<dgm:relIds xmlns:dgm="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:dm="%s" r:lo="%s" r:qs="%s" r:cs="%s"/>'
                       % (D, r_dm, r_lo, r_qs, r_cs), "Diagram " + n, alt)


def add_chart(doc, p, kind="col", alt=""):
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    n = next_id()
    chart = ('<c:chartSpace xmlns:c="%s"><c:chart><c:plotArea><c:barChart><c:barDir val="%s"/><c:grouping val="clustered"/></c:barChart></c:plotArea></c:chart></c:chartSpace>' % (C, kind))
    rid = add_part(doc, "/word/charts/chart%s.xml" % n, "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", chart,
                   "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart")
    add_graphic_inline(p, C, '<c:chart xmlns:c="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="%s"/>' % (C, rid), "Chart " + n, alt)


def add_model3d(p, alt=""):
    add_graphic_inline(p, "http://schemas.microsoft.com/office/drawing/2017/model3d",
                       '<am3d:model3d xmlns:am3d="http://schemas.microsoft.com/office/drawing/2017/model3d"><am3d:trans><am3d:rot ax="1200000" ay="-2400000" az="0"/></am3d:trans></am3d:model3d>', "3D Model", alt)


def make_icon(doc, index):
    """Mark the nth picture as an icon (Insert > Icons adds SVG pictures)."""
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    pic = list(doc.element.body.iter("{http://schemas.openxmlformats.org/drawingml/2006/picture}pic"))[index]
    blip = next(pic.iter("{%s}blip" % A))
    blip.append(parse_xml('<a:extLst xmlns:a="%s"><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}"><asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
                          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="%s"/></a:ext></a:extLst>'
                          % (A, blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"))))


# --------------------------------------------------------------------------- citations, themes, merge
def sources(doc, count, style="APA", version="7"):
    """The citation manager's source list (a customXml part) with `count` practice sources."""
    B = "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
    items = "".join('<b:Source><b:Tag>Src%d</b:Tag><b:SourceType>JournalArticle</b:SourceType><b:Title>Practice source %d</b:Title></b:Source>' % (i, i) for i in range(1, count + 1))
    xml = '<b:Sources xmlns:b="%s" xmlns="%s" SelectedStyle="\\APASixthEditionOfficeOnline.xsl" StyleName="%s" Version="%s">%s</b:Sources>' % (B, B, style, version, items)
    part = Part(PackURI("/customXml/item9.xml"), "application/xml", xml.encode("utf-8"), doc.part.package)
    doc.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml")


def theme_colors(doc, name, accent1):
    """Rename the theme's color set and change Accent 1 (Design > Colors > Customize Colors)."""
    from lxml import etree
    for rel in doc.part.rels.values():
        if rel.reltype == RT.THEME:
            part = rel.target_part
            root = etree.fromstring(part.blob)
            A = "http://schemas.openxmlformats.org/drawingml/2006/main"
            scheme = root.find(".//{%s}clrScheme" % A)
            scheme.set("name", name)
            acc = scheme.find("{%s}accent1" % A)
            for c in list(acc):
                acc.remove(c)
            etree.SubElement(acc, "{%s}srgbClr" % A).set("val", accent1)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def mail_merge(doc, kind="formLetters", source=True):
    src = '<w:dataSource r:id="rIdMergeSrc"/>' if False else ""
    q = '<w:query w:val="SELECT * FROM `Employers`"/>' if source else ""
    setting(doc, '<w:mailMerge %s><w:mainDocumentType w:val="%s"/><w:linkToQuery/><w:dataType w:val="textFile"/><w:connectString w:val=""/>%s%s</w:mailMerge>' % (NSW, kind, q, src))


def content_control(p, placeholder_text, showing_placeholder=True):
    """A plain-text content control (the [Your name] boxes in templates)."""
    show = "<w:showingPlcHdr/>" if showing_placeholder else ""
    raw_run(p, '<w:sdt %s><w:sdtPr><w:alias w:val="Text"/><w:tag w:val="field"/>%s<w:text/></w:sdtPr><w:sdtContent><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:sdtContent></w:sdt>' % (NSW, show, esc(placeholder_text)))


def save(doc, folder, name, root):
    path = folder / name
    folder.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    print("wrote", path.relative_to(root))
